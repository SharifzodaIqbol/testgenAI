"""
CRUD для тест-кейсов + экспорт в Excel и JSON.
"""
from __future__ import annotations

import io
import json
from typing import List

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.security import get_current_user_id
from app.db.base import Document, Project, TestCase, TestSuite
from app.db.session import get_db
from app.schemas import TestCaseCreate, TestCaseRead, TestCaseUpdate, TestSuiteRead
from urllib.parse import quote

router = APIRouter()


# ── Suites ─────────────────────────────────────────────────────────────────────

@router.get("/suites/project/{project_id}", response_model=List[TestSuiteRead])
async def list_suites(
    project_id: int,
    user_id: int = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    project = await db.get(Project, project_id)
    if not project or project.owner_id != user_id:
        raise HTTPException(404, "Проект не найден")

    suites = await db.scalars(
        select(TestSuite)
        .where(TestSuite.project_id == project_id)
        .options(selectinload(TestSuite.test_cases))
        .order_by(TestSuite.created_at.desc())
    )
    return suites.all()


@router.get("/suites/{suite_id}", response_model=TestSuiteRead)
async def get_suite(
    suite_id: int,
    user_id: int = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    suite = await db.scalar(
        select(TestSuite)
        .where(TestSuite.id == suite_id)
        .options(selectinload(TestSuite.test_cases))
    )
    if not suite:
        raise HTTPException(404, "Suite не найден")

    project = await db.get(Project, suite.project_id)
    if not project or project.owner_id != user_id:
        raise HTTPException(403, "Нет прав")

    return suite


@router.delete("/suites/{suite_id}", status_code=204)
async def delete_suite(
    suite_id: int,
    user_id: int = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    suite = await db.get(TestSuite, suite_id)
    if not suite:
        raise HTTPException(404, "Suite не найден")
    project = await db.get(Project, suite.project_id)
    if not project or project.owner_id != user_id:
        raise HTTPException(403, "Нет прав")
    await db.delete(suite)


# ── Test Cases ─────────────────────────────────────────────────────────────────

@router.post("/suites/{suite_id}/cases", response_model=TestCaseRead, status_code=201)
async def create_test_case(
    suite_id: int,
    body: TestCaseCreate,
    user_id: int = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    suite = await db.get(TestSuite, suite_id)
    if not suite:
        raise HTTPException(404, "Suite не найден")
    project = await db.get(Project, suite.project_id)
    if not project or project.owner_id != user_id:
        raise HTTPException(403, "Нет прав")

    tc = TestCase(
        suite_id=suite_id,
        **body.model_dump(),
    )
    db.add(tc)
    await db.flush()
    return tc


@router.patch("/{case_id}", response_model=TestCaseRead)
async def update_test_case(
    case_id: int,
    body: TestCaseUpdate,
    user_id: int = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    tc = await db.get(TestCase, case_id)
    if not tc:
        raise HTTPException(404, "Тест-кейс не найден")
    await _check_tc_rights(tc, user_id, db)

    for field, value in body.model_dump(exclude_unset=True).items():
        setattr(tc, field, value)
    await db.flush()
    return tc


@router.delete("/{case_id}", status_code=204)
async def delete_test_case(
    case_id: int,
    user_id: int = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    tc = await db.get(TestCase, case_id)
    if not tc:
        raise HTTPException(404, "Тест-кейс не найден")
    await _check_tc_rights(tc, user_id, db)
    await db.delete(tc)


# ── Export ─────────────────────────────────────────────────────────────────────

@router.get("/suites/{suite_id}/export/json")
async def export_json(
    suite_id: int,
    user_id: int = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    suite = await _get_suite_with_cases(suite_id, user_id, db)
    data = {
        "suite_name": suite.name,
        "test_cases": [
            {
                "id": tc.id,
                "title": tc.title,
                "description": tc.description,
                "preconditions": tc.preconditions,
                "steps": tc.steps,
                "expected_result": tc.expected_result,
                "priority": tc.priority.value,
                "case_type": tc.case_type.value,
                "tags": tc.tags,
                "confidence_score": tc.confidence_score,
            }
            for tc in suite.test_cases
        ],
    }
    
    # Превращаем JSON в байтовую строку UTF-8
    content = json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
    
    # Безопасно кодируем имя файла для HTTP заголовков (RFC 5987)
    safe_name = quote(suite.name.replace(" ", "_"))
    
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/json",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{safe_name}.json"
        },
    )


@router.get("/suites/{suite_id}/export/excel")
async def export_excel(
    suite_id: int,
    user_id: int = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    try:
        import openpyxl  # type: ignore
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(500, "openpyxl не установлен")

    suite = await _get_suite_with_cases(suite_id, user_id, db)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = suite.name[:31]

    headers_list = [
        "ID", "Название", "Описание", "Предусловия", "Шаги",
        "Ожидаемый результат", "Приоритет", "Тип", "Теги", "Confidence",
    ]
    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers_list, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row, tc in enumerate(suite.test_cases, 2):
        steps_text = "\n".join(
            f"{i+1}. {s['step']}\n   → {s['expected']}"
            for i, s in enumerate(tc.steps)
        )
        ws.append([
            tc.id, tc.title, tc.description, tc.preconditions,
            steps_text, tc.expected_result,
            tc.priority.value, tc.case_type.value,
            ", ".join(tc.tags), round(tc.confidence_score, 2),
        ])

    # Ширина колонок
    col_widths = [8, 40, 40, 30, 60, 40, 12, 15, 20, 12]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # Безопасно кодируем имя файла для HTTP заголовков (RFC 5987)
    safe_name = quote(suite.name.replace(" ", "_"))
    
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{safe_name}.xlsx"
        },
    )


# ── Helpers ────────────────────────────────────────────────────────────────────

async def _check_tc_rights(tc: TestCase, user_id: int, db: AsyncSession):
    suite = await db.get(TestSuite, tc.suite_id)
    project = await db.get(Project, suite.project_id)
    if not project or project.owner_id != user_id:
        raise HTTPException(403, "Нет прав")


async def _get_suite_with_cases(suite_id: int, user_id: int, db: AsyncSession) -> TestSuite:
    suite = await db.scalar(
        select(TestSuite)
        .where(TestSuite.id == suite_id)
        .options(selectinload(TestSuite.test_cases))
    )
    if not suite:
        raise HTTPException(404, "Suite не найден")
    project = await db.get(Project, suite.project_id)
    if not project or int(project.owner_id) != int(user_id):
        raise HTTPException(403, "Нет прав")
    return suite